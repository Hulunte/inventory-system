"""Integration and resource checks for the Phase 2 ticket workflow."""
import io
from datetime import datetime, timezone
from decimal import Decimal
from unittest.mock import patch

from openpyxl import load_workbook

from app.extensions import db
from app.models.harvest_entry import HarvestEntry
from app.models.product import Product
from tests.conftest import make_worker_with_assignment


DAY = "2026-06-15"


def _entry(db_session, *, amount="25.00", voided=False):
    worker, assignment = make_worker_with_assignment(
        db_session, slot_number=149, person_name="Trabajador Fase Dos"
    )
    product = Product(name="Producto Fase Dos", rate_per_kg=Decimal("5.00"), active=True)
    db_session.add(product)
    db_session.flush()
    entry = HarvestEntry(
        worker_id=worker.id,
        worker_assignment_id=assignment.id,
        product_id=product.id if amount is not None else None,
        weight_kg=Decimal("5.000"),
        amount_mxn=Decimal(amount) if amount is not None else None,
        created_at=datetime(2026, 6, 15, 12, tzinfo=timezone.utc),
        voided=voided,
        worker_name_snapshot=assignment.person_name,
        worker_barcode_snapshot=worker.barcode,
        worker_slot_number_snapshot=worker.slot_number,
        product_name_snapshot=product.name if amount is not None else None,
        rate_per_kg_snapshot=product.rate_per_kg if amount is not None else None,
    )
    db_session.add(entry)
    db_session.flush()
    return worker, assignment, entry


def test_reception_has_visible_go_to_tickets_button(client):
    html = client.get("/").get_data(as_text=True)
    assert 'class="scanner-zone__tickets"' in html
    assert 'href="/tickets#quick-scan-section"' in html
    assert 'aria-label="Ir a impresión rápida de tickets"' in html
    assert "Ir a tickets" in html


def test_reception_ticket_button_is_responsive_and_does_not_wrap_scanner(client):
    css = client.get("/static/css/reception.css").get_data(as_text=True)
    assert ".scanner-zone__heading" in css
    assert "justify-content: space-between" in css
    assert "@media (max-width: 520px)" in css
    html = client.get("/").get_data(as_text=True)
    assert html.index("scanner-zone__tickets") < html.index('id="barcode"')


def test_quick_scanner_mode_and_all_print_modes_are_rendered(admin_client):
    html = admin_client.get("/tickets").get_data(as_text=True)
    assert 'id="quick-scan-enabled"' in html
    assert 'id="quick-scan-input"' in html
    assert 'id="quick-scan-date"' in html
    assert "Impresi&oacute;n r&aacute;pida por esc&aacute;ner" in html
    assert 'id="print-selected-btn"' in html
    assert 'id="print-all-auto-btn"' in html
    assert 'id="print-all-confirm-btn"' in html


def test_ticket_javascript_supports_scanner_toggle_queue_controls_and_dedup(client):
    js = client.get("/static/js/tickets.js").get_data(as_text=True)
    assert "processQuickScan" in js
    assert "quickPrintedTickets" in js
    assert "has_incomplete_amounts" in js
    assert "confirm-incomplete" not in js
    assert "confirm-skip-btn" in js
    assert "confirm-stop-btn" in js
    assert "error-retry-btn" in js
    assert "selectedAssignmentId === assignmentId" in js
    assert 'window.location.hash === "#quick-scan-section"' in js
    assert "getOperationalDate" in js
    assert "Configure primero una impresora" in js


def test_ticket_page_supplies_timezone_aware_operational_date(admin_client, monkeypatch):
    from datetime import date

    monkeypatch.setattr("app.routes.tickets._operational_today", lambda: date(2026, 9, 8))
    html = admin_client.get("/tickets").get_data(as_text=True)
    assert 'operationalToday: "2026-09-08"' in html


def test_quick_scan_date_falls_back_and_tracks_manual_date(client):
    js = client.get("/static/js/tickets.js").get_data(as_text=True)
    assert "dateInput.value || configuredToday" in js
    assert "quickScanDate.textContent" in js
    assert 'dateInput.addEventListener("change", getOperationalDate)' in js


def test_quick_scan_checks_printer_before_printing(client):
    js = client.get("/static/js/tickets.js").get_data(as_text=True)
    printer_check = js.index("if (!printerName)", js.index("async function processQuickScan"))
    print_call = js.index("const result = await sendPrintRequest", printer_check)
    assert printer_check < print_call
    assert "quickScanBusy" in js


def test_quick_scanner_lookup_has_no_movements(admin_client):
    response = admin_client.get(f"/api/tickets/daily?date={DAY}&q=TRB999999")
    assert response.status_code == 200
    assert response.get_json()["tickets"] == []


@patch("app.routes.tickets.print_raw")
@patch("app.routes.tickets.validate_printer_name", return_value=True)
@patch("app.routes.tickets.is_available", return_value=True)
def test_scanner_ticket_print_requires_incomplete_confirmation(
    available, valid_printer, print_raw, admin_client, db_session
):
    _, assignment, _ = _entry(db_session, amount=None)
    db_session.commit()
    csrf = admin_client.get("/api/admin/session").get_json()["csrf_token"]
    body = {"date": DAY, "worker_assignment_id": assignment.id, "printer_name": "Mock"}

    pending = admin_client.post(
        "/api/tickets/print", json=body, headers={"X-CSRF-Token": csrf}
    )
    assert pending.status_code == 409
    assert pending.get_json()["has_incomplete_amounts"] is True
    print_raw.assert_not_called()

    body["confirm_incomplete_amounts"] = True
    printed = admin_client.post(
        "/api/tickets/print", json=body, headers={"X-CSRF-Token": csrf}
    )
    assert printed.status_code == 200
    print_raw.assert_called_once()


def test_report_summary_amount_total_and_voided_exclusion(admin_client, db_session):
    worker, assignment, active = _entry(db_session, amount="25.00")
    db_session.add(HarvestEntry(
        worker_id=worker.id,
        worker_assignment_id=assignment.id,
        product_id=active.product_id,
        weight_kg=Decimal("2.000"),
        amount_mxn=Decimal("10.00"),
        created_at=datetime(2026, 6, 15, 13, tzinfo=timezone.utc),
        voided=True,
        voided_at=datetime(2026, 6, 15, 14, tzinfo=timezone.utc),
        void_reason="Anulación de prueba",
        worker_name_snapshot=active.worker_name_snapshot,
        worker_barcode_snapshot=active.worker_barcode_snapshot,
        worker_slot_number_snapshot=active.worker_slot_number_snapshot,
        product_name_snapshot=active.product_name_snapshot,
        rate_per_kg_snapshot=active.rate_per_kg_snapshot,
    ))
    db_session.commit()

    response = admin_client.get(
        f"/api/reports/harvest/export?start_date={DAY}&end_date={DAY}"
    )
    workbook = load_workbook(io.BytesIO(response.data))
    summary = workbook["Resumen"]
    assert summary.cell(2, 6).value == 25
    assert summary.cell(2, 6).number_format == '$#,##0.00'
    assert summary.cell(summary.max_row, 6).value == 25
    assert summary.cell(2, 7).value == 1


def test_report_summary_marks_incomplete_amounts(admin_client, db_session):
    _entry(db_session, amount=None)
    db_session.commit()
    response = admin_client.get(
        f"/api/reports/harvest/export?start_date={DAY}&end_date={DAY}"
    )
    summary = load_workbook(io.BytesIO(response.data))["Resumen"]
    assert summary.cell(2, 6).value == "N/D"
    assert summary.cell(summary.max_row, 6).value == "N/D"
